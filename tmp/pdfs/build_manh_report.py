from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

BASE = Path('output/baitap1_manh').resolve()
R = json.loads((BASE / 'ket_qua/results.json').read_text())
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.0)
sec.footer_distance = Cm(0.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    rf = st.element.get_or_add_rPr().find(qn('w:rFonts'))
    if rf is not None:
        for key in list(rf.attrib):
            if key.endswith('Theme'):
                del rf.attrib[key]
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(21)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.styles['Caption'].font.size = Pt(10.5)
doc.styles['Caption'].font.italic = True
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)
doc.core_properties.author = 'Lương Đức Mạnh'
doc.core_properties.title = 'Phân tích và ra quyết định Bayes cho lỗi máy chủ hiếm gặp'


def p(text, bold=False):
    x = doc.add_paragraph()
    x.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    x.add_run(text).bold = bold
    return x


def h(text, level=1):
    doc.add_heading(text, level=level)


def page(title):
    para = doc.add_heading(title, level=1)
    para.paragraph_format.page_break_before = True


def cap(text):
    x = doc.add_paragraph(text, 'Caption')
    x.paragraph_format.keep_with_next = True


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = str(text)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:' + side)
        for k,v in [('val','single'),('sz','5'),('color','D9D9D9')]:
            el.set(qn('w:' + k), v)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if i == 0: pr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            if widths: c.width = Cm(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'E6EDF2' if i == 0 else 'FFFFFF')
            tcpr.append(sh)
            mar = OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), '70'); el.set(qn('w:type'), 'dxa'); mar.append(el)
            tcpr.append(mar)
            for para in c.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.space_before = Pt(1)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    run.font.size = Pt(10.5)
                    run.bold = i == 0
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer.add_run(' ').font.size = Pt(1)
    return t


def mathrun(text):
    r = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = text; r.append(t); return r


def mathparts(parent, text):
    import re
    pos = 0
    for match in re.finditer(r'([CLp])_(FP|FN|Test|đổi)', text):
        if match.start() > pos:
            parent.append(mathrun(text[pos:match.start()]))
        sub = OxmlElement('m:sSub')
        base = OxmlElement('m:e'); base.append(mathrun(match.group(1)))
        index = OxmlElement('m:sub'); index.append(mathrun(match.group(2)))
        sub.extend([base,index]); parent.append(sub)
        pos = match.end()
    if pos < len(text):
        parent.append(mathrun(text[pos:]))


def eq(*parts):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    om = OxmlElement('m:oMath')
    for part in parts:
        if isinstance(part, tuple):
            frac = OxmlElement('m:f')
            for tag, text in zip(['num','den'], part):
                e = OxmlElement('m:' + tag); mathparts(e,text); frac.append(e)
            om.append(frac)
        else:
            mathparts(om,part)
    para._p.append(om)


def fig(name, caption):
    cap(caption)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(BASE / 'ket_qua' / name), width=Cm(14.0 if name == 'likelihood.png' else 15.4))


def fmt(x, n=4):
    return f'{x:.{n}f}'.replace('.', ',')


# Trang 1
for text in ['BÀI TẬP 1', 'HỌC PHẦN NHẬN DẠNG MẪU']:
    x = doc.add_paragraph(text)
    x.alignment = WD_ALIGN_PARAGRAPH.CENTER
    x.runs[0].bold = True
doc.add_paragraph()
x = doc.add_paragraph('Phân tích và ra quyết định Bayes cho lỗi máy chủ hiếm gặp', 'Title')
x.alignment = WD_ALIGN_PARAGRAPH.CENTER
x.paragraph_format.space_after = Pt(22)
x = doc.add_paragraph('Phát hiện lỗi nghiêm trọng trong trung tâm dữ liệu', 'Subtitle')
x.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()
for text in ['Sinh viên  Lương Đức Mạnh', 'Mã học viên  B26CHKH068', 'Lớp  M26CQKH03-B', 'Năm học  2026 - 2027']:
    x = doc.add_paragraph(text); x.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()
h('Tóm tắt', 2)
p('Báo cáo xây dựng bộ phân loại Bayes để phát hiện lỗi máy chủ từ số chỉ báo bất thường. Dữ liệu mô phỏng gồm 4.800 mẫu Train và 1.200 mẫu Test, với tỷ lệ lỗi 4,5%. Prior và likelihood được ước lượng từ Train; posterior được dùng để so sánh quy tắc MAP với quyết định tối thiểu hóa rủi ro có điều kiện.')
p('Khi chi phí bỏ sót bằng 20 lần chi phí cảnh báo nhầm, quy tắc theo chi phí giảm số lỗi bỏ sót từ 25 xuống 5 và tổng chi phí trên Test từ 506 xuống 266, đồng thời tăng số cảnh báo nhầm. Kết quả minh họa vai trò của tỷ lệ nền và hậu quả sai lầm trong quyết định. Toàn bộ dữ liệu và phép tính có thể tái tạo bằng chương trình đi kèm.')
p('Phạm vi dữ liệu: các tần số được thiết kế cho thực nghiệm mô phỏng; kết quả không phải bằng chứng về hiệu quả vận hành trên hệ thống máy chủ thực.', bold=True)

# Trang 2
page('1 Bài toán và dữ liệu mô phỏng')
p('Trong trung tâm dữ liệu, lỗi nghiêm trọng như dừng dịch vụ hoặc lỗi phần cứng có thể xuất hiện với tỷ lệ thấp. Bài toán là dự đoán trạng thái lỗi hiện tại từ các chỉ báo giám sát, sau đó quyết định có phát cảnh báo để kiểm tra hay không. Mô hình này không dự báo thời điểm xảy ra lỗi trong tương lai.')
h('1 1 Biểu diễn mẫu', 2)
p('Mỗi mẫu là cặp (x, y), trong đó y = 1 biểu thị Failure và y = 0 biểu thị Normal. Đặc trưng x là số chỉ báo bất thường được gộp thành năm bin: 0, 1, 2, 3 và ≥4. Trong mã, giá trị 4 đại diện cho toàn bộ bin ≥4, không phải chỉ đúng bốn chỉ báo.')
p('Trong ứng dụng thực, các chỉ báo có thể được rút ra từ CPU, nhiệt độ, RAM, độ trễ I/O và cảnh báo phần cứng. Bài này mô phỏng trực tiếp đặc trưng sau khi gộp bin; không tạo hoặc sử dụng số đo telemetry thô. Do đó, chương trình đánh giá bước phân loại từ x, chưa đánh giá bước trích xuất chỉ báo từ cảm biến.')
h('1 2 Thiết kế hai tập dữ liệu', 2)
cap('Bảng 1. Tần số mô phỏng theo lớp và mức x')
table(['x','Train Normal','Train Failure','Test Normal','Test Failure'], [
    ['0',3900,20,980,5],['1',600,35,140,8],['2',70,50,20,12],['3',10,55,4,14],['≥4',4,56,2,15],['Tổng',4584,216,1146,54]], [1.3,3.85,3.85,3.85,3.85])
p('Hai tập được ấn định bằng bảng tần số trên rồi mở rộng thành từng dòng dữ liệu. Đây là thiết kế mô phỏng theo tần số cố định, không phải kết quả chia ngẫu nhiên một bộ dữ liệu thu thập thực tế. Không cần seed ngẫu nhiên; chạy lại cùng mã luôn thu được cùng dữ liệu.')
p('Tổng số mẫu là 6.000. Train gồm 4.800 mẫu, có 216 mẫu lỗi; Test gồm 1.200 mẫu, có 54 mẫu lỗi. Hai tập đều có tỷ lệ lỗi 4,5% theo thiết kế. Chương trình chỉ học tham số trên Train và sử dụng Test để tính chỉ số đánh giá. Nhãn Test không được dùng để ước lượng prior, likelihood hoặc chọn chi phí.')

# Trang 3
page('2 Ước lượng prior và likelihood')
h('2 1 Prior phản ánh tỷ lệ lỗi nền', 2)
eq('P(F) = ', ('216','4800'), ' = 0,045;    P(N) = 0,955')
p('F là Failure, N là Normal. Prior thể hiện xác suất trước khi biết x: một mẫu trong phân bố Train có xác suất lỗi 4,5%. Sự kiện lỗi vì vậy là lớp thiểu số trong thực nghiệm này.')
h('2 2 Likelihood từ tần số theo lớp', 2)
p('Likelihood P(x | c) mô tả khả năng quan sát mức x khi đã biết lớp c. Với K = 5 bin và hệ số làm mượt α = 1, ước lượng được tính như sau:')
eq('P(x = k | c) = ', ('n(c,k) + α','n(c) + Kα'))
p('Trong đó n(c,k) là số mẫu lớp c thuộc bin k, còn n(c) là tổng số mẫu của lớp đó. Các ô hiện tại đều khác 0; Laplace smoothing vẫn được sử dụng để giảm mức cực đoan của ước lượng ở những ô ít mẫu và xử lý thống nhất trường hợp có ô bằng 0.')
cap('Bảng 2. Likelihood sau làm mượt trên Train')
table(['x','P(x | N)','P(x | F)'], [[str(i) if i<4 else '≥4',fmt(R['likelihood'][0][i]),fmt(R['likelihood'][1][i])] for i in range(5)], [3,6.85,6.85])
fig('likelihood.png','Hình 1. Phân bố likelihood theo từng lớp')
p('Normal có khoảng 85,01% khối xác suất tại x = 0; Failure có khoảng 74,21% tại x ≥ 2. Likelihood mỗi lớp cộng bằng 1. Chỉ likelihood được làm mượt; prior vẫn dùng tần suất lớp.')

# Trang 4
page('3 Posterior và quy tắc quyết định')
p('Với một quan sát mới thuộc bin x, định lý Bayes kết hợp bằng chứng quan sát với tỷ lệ lỗi nền:')
eq('P(F | x) = ', ('P(x | F)P(F)','P(x | F)P(F) + P(x | N)P(N)'))
p('Đặt q = P(F | x), khi đó P(N | x) = 1 − q. MAP chọn lớp có posterior lớn nhất, tương đương chọn Failure khi q ≥ 0,5. Trường hợp hòa được quy ước chọn cảnh báo trong chương trình.')
h('3 1 Quyết định có xét hậu quả sai lầm', 2)
cap('Bảng 3. Ma trận chi phí giả định theo quyết định và trạng thái thật')
table(['Quyết định','Thật Normal','Thật Failure'],[['Cảnh báo', 'C_FP = 1', '0'],['Không cảnh báo','0','C_FN = 20']], [7,4.85,4.85])
p('Một đơn vị chi phí là thang đo tương đối, không quy đổi thành tiền. Chi phí quyết định đúng được giả định bằng 0. Cảnh báo nhầm gây thêm việc kiểm tra; bỏ sót lỗi có thể gây gián đoạn dịch vụ. Tỷ lệ 20 được chọn để khảo sát, chưa được hiệu chỉnh bằng số liệu thiệt hại thực tế.')
eq('R(cảnh báo | x) = C_FP(1 − q);    R(không cảnh báo | x) = C_FN q')
eq('q ≥ t* = ', ('C_FP','C_FP + C_FN'), ' = ', ('1','21'), ' ≈ 0,047619')
p('MAP tương ứng với chi phí hai loại sai lầm bằng nhau trong mô hình mất mát 0–1. Khi chi phí bất đối xứng, chọn hành động có rủi ro có điều kiện thấp hơn; nếu hai rủi ro bằng nhau, cả hai hành động cùng tối ưu và mã chọn cảnh báo.')
cap('Bảng 4. Posterior và quyết định theo từng bin')
table(['x','P(F | x)','P(N | x)','MAP','Theo chi phí'], [[str(i) if i<4 else '≥4',fmt(q),fmt(1-q),'Lỗi' if q>=.5 else 'Bình thường','Cảnh báo' if q>=1/21 else 'Không'] for i,q in enumerate(R['posterior'])], [1.1,3.15,3.15,4.5,4.8])
p('Ví dụ x = 2: q ≈ 0,412740. MAP chọn Normal, nhưng rủi ro cảnh báo ≈ 0,587260 nhỏ hơn rủi ro không cảnh báo ≈ 8,254807 nên quy tắc theo chi phí phát cảnh báo. Đó là quyết định kiểm tra dựa trên chi phí, không phải khẳng định máy chắc chắn hỏng.')

# Trang 5
page('4 Đánh giá trên tập Test')
p('TP là mẫu lỗi được cảnh báo đúng; FP là mẫu bình thường bị cảnh báo nhầm; FN là mẫu lỗi bị bỏ sót; TN là mẫu bình thường được dự đoán đúng. Các phép đếm sử dụng đủ 1.200 dòng Test được tạo từ Bảng 1.')
eq('Accuracy = ', ('TP + TN','TP + FP + FN + TN'))
eq('Precision = ', ('TP','TP + FP'), ';    Recall = ', ('TP','TP + FN'))
eq('L_Test = C_FP × FP + C_FN × FN')
cap('Bảng 5. Kết quả Test với cùng chi phí C_FP = 1 và C_FN = 20')
table(['Chỉ tiêu','Luôn Normal','MAP','Theo chi phí'], [
    ['Ngưỡng posterior','Không cảnh báo','0,5','1/21'],['TP',0,29,49],['FP',0,6,166],['FN',54,25,5],['TN',1146,1140,980],
    ['Accuracy','95,50%','97,42%','85,75%'],['Precision','Không xác định','82,86%','22,79%'],['Recall','0,00%','53,70%','90,74%'],['Tổng chi phí',1080,506,266],['Chi phí mỗi mẫu','0,9000','0,4217','0,2217']], [5.0,3.9,3.9,3.9])
p('MAP cảnh báo ở x ≥ 3 nên TP = 14 + 15 = 29 và FP = 4 + 2 = 6. Quy tắc theo chi phí cảnh báo ở x ≥ 1 nên TP = 8 + 12 + 14 + 15 = 49 và FP = 140 + 20 + 4 + 2 = 166. Các số FN, TN là phần còn lại của từng lớp.')
p('Tổng chi phí giảm 240 đơn vị, tương đương 47,43% so với MAP. Tuy nhiên, số cảnh báo tăng từ 35 lên 215 và Precision giảm mạnh. Hệ thống vận hành cần cân nhắc khả năng xử lý số cảnh báo này.')
p('Cách luôn chọn Normal vẫn đạt Accuracy 95,50% nhưng bỏ sót toàn bộ 54 mẫu lỗi. Vì vậy, Accuracy riêng lẻ không phản ánh đầy đủ mục tiêu phát hiện sự kiện hiếm. Precision của cách này không xác định vì không có dự đoán dương tính; hàm đánh giá dùng giá trị quy ước 0 nếu mẫu số bằng 0 để tránh lỗi chia.')
p('L_Test là chi phí thực nghiệm trên nhãn Test. L_Test/1.200 là chi phí trung bình thực nghiệm; chúng khác với rủi ro có điều kiện tính từ posterior cho một quan sát. Tối ưu rủi ro theo mô hình không bảo đảm tối ưu trên mọi tập dữ liệu hữu hạn hoặc trên dữ liệu thực.')

# Trang 6
page('5 Ảnh hưởng của chi phí sai lầm')
p('Giữ nguyên dữ liệu và posterior, đặt C_FP = 1 rồi thay đổi r = C_FN/C_FP. Ngưỡng cảnh báo là t = 1/(1+r). Mỗi hàng dưới đây đánh giá hai quy tắc với cùng chi phí của chính kịch bản đó.')
cap('Bảng 6. Khảo sát chi phí trên cùng tập Test')
table(['r','Ngưỡng','TP','FP','FN','Recall %','L theo chi phí','L MAP'], [[row['ratio'],fmt(row['threshold']),row['TP'],row['FP'],row['FN'],fmt(row['recall']*100,2),row['cost'],row['MAP_cost']] for row in R['cost_sweep']], [1,2.15,1.25,1.35,1.25,2.3,3.65,3.6])
fig('cost.png','Hình 2. Precision và Recall tại các kịch bản chi phí được khảo sát')
p('Khi r tăng từ 1 lên 2, ngưỡng giảm từ 0,5 xuống 1/3 và bin x = 2 được cảnh báo thêm. Khi r tăng đến 20, bin x = 1 được cảnh báo thêm. Bin x = 0 vẫn chưa được cảnh báo ở r = 100 vì posterior khoảng 0,00524 thấp hơn ngưỡng 1/101.')
p('Các kịch bản r = 2, 5, 10 có cùng tập dự đoán; tương tự với r = 20, 50, 100. Vì đặc trưng chỉ có năm bin, chỉ có năm mức posterior nên Recall không tăng liên tục mỗi khi thay đổi chi phí. Đường nối trong Hình 2 chỉ giúp theo dõi các kịch bản rời rạc.')
p('Tổng chi phí giữa các hàng không có cùng thang phạt bỏ sót nên không thể dùng để chọn “kịch bản tốt nhất”. So sánh phù hợp là giữa hai quy tắc trong cùng một hàng. Trên dữ liệu mô phỏng này, các kịch bản được khảo sát đều cho chi phí của quy tắc theo rủi ro không lớn hơn MAP.')

# Trang 7
page('6 Ảnh hưởng của tỷ lệ lỗi nền')
p('Giữ likelihood đã học trên Train, thay prior p = P(F) và tính lại cả tử số lẫn mẫu số của posterior. Xét cùng một quan sát x = 2; đây là khảo sát quyết định cho một quan sát, không phải kết quả đánh giá một tập Test mới.')
cap('Bảng 7. Posterior tại x = 2 theo prior giả định')
table(['Prior P(F)','Posterior P(F | x = 2)','MAP'], [[fmt(row['prior'],3),fmt(row['posterior']),'Failure' if row['MAP'] else 'Normal'] for row in R['prior_sweep']], [4,7,5.7])
fig('prior.png','Hình 3. Posterior tại x = 2 khi giữ nguyên likelihood và thay prior')
eq('p_đổi = ', ('P(x = 2 | N)','P(x = 2 | N) + P(x = 2 | F)'), ' ≈ 0,0628319')
p('Tại prior khoảng 6,2832%, posterior của hai lớp bằng nhau. Với quy ước hòa chọn Failure, MAP chuyển sang Failure từ điểm này; dưới điểm này MAP chọn Normal. Chẳng hạn, prior 4,5% cho posterior 41,2740%, còn prior 10% cho posterior 62,3675%.')
p('Việc chỉ thay prior hợp lý khi phân bố đặc trưng trong từng lớp giữ ổn định. Nếu nhiệt độ vận hành, phần cứng hoặc cách phát sinh cảnh báo làm P(x | lớp) thay đổi, cần ước lượng lại likelihood; không thể mặc nhiên tái sử dụng mô hình cũ.')

# Trang 8
page('7 Giới hạn và kết luận')
h('7 1 Giới hạn của thực nghiệm', 2)
p('Dữ liệu là các tần số thiết kế sẵn để minh họa Bayes. Hai tập Train và Test riêng nhau về vai trò trong chương trình, nhưng không phải các mẫu ngẫu nhiên độc lập thu từ môi trường vận hành. Vì vậy, các chỉ số phản ánh kịch bản mô phỏng cụ thể và không chứng minh khả năng khái quát trên máy chủ thực.')
p('Gộp nhiều cảnh báo thành một số đếm làm mất thông tin về loại cảnh báo và mức độ nghiêm trọng. Hai máy có cùng x có thể mang các dấu hiệu rất khác nhau. Bài này dùng một đặc trưng phân loại rời rạc, không cần giả định độc lập có điều kiện giữa nhiều đặc trưng như Naive Bayes.')
p('Một số bin Normal chỉ có ít mẫu; ước lượng likelihood có thể nhạy với α và cách gộp bin. Báo cáo cố định α = 1 và chưa khảo sát độ bất định thống kê hoặc độ hiệu chuẩn posterior. Chi phí 1 và 20 cũng chỉ là giả định cho bài tập.')
p('Nếu triển khai thực tế, cần xây dựng quy tắc trích xuất chỉ báo nhất quán, thu nhãn lỗi tin cậy, tách dữ liệu theo thời gian hoặc theo máy để hạn chế rò rỉ, và đánh giá trên tập độc lập. Chi phí và năng lực kiểm tra cảnh báo cần được xác định trước khi lựa chọn quy tắc vận hành.')
h('7 2 Kết luận', 2)
p('Mô hình đã biểu diễn mẫu bằng x_bin, ước lượng prior và likelihood từ dữ liệu Train, tính posterior cho quan sát mới và so sánh hai quy tắc quyết định. Với sự kiện lỗi có prior thấp, bằng chứng bất thường chưa nhất thiết đủ để MAP chọn lớp lỗi.')
p('Khi C_FN = 20 và C_FP = 1, ngưỡng cảnh báo giảm từ 0,5 xuống 1/21. Trên Test mô phỏng, số bỏ sót giảm từ 25 xuống 5, Recall tăng từ 53,70% lên 90,74% và tổng chi phí giảm từ 506 xuống 266. Sự cải thiện về chi phí đi kèm 160 cảnh báo nhầm tăng thêm.')
p('Khảo sát cho thấy thay prior hoặc chi phí có thể đảo quyết định cho cùng quan sát. Kết luận luôn phụ thuộc vào giả định về likelihood, tỷ lệ nền và hậu quả sai lầm. Toàn bộ số liệu trong báo cáo được sinh và tính lại bằng mã ở phụ lục, với các bảng dữ liệu đầu vào được công bố tường minh.')
h('Nguồn dữ liệu và bài toán', 2)
p('Yêu cầu bài toán: Baitap1.pdf, bài tập Phân tích và ra quyết định Bayes cho một sự kiện hiếm. Nguồn số liệu thực nghiệm: bảng tần số mô phỏng tại Bảng 1 của báo cáo này. Các biểu đồ và kết quả là đầu ra tính toán từ bảng đó; không sử dụng dữ liệu giám sát thực tế.')

# Trang 9
page('8 Chương trình và cách tái tạo kết quả')
p('Chương trình bayes_server.py dùng thư viện chuẩn của Python 3. Phụ lục A chứa toàn bộ mã tính toán; tệp cùng tên được cung cấp để chạy trực tiếp, tránh lỗi khi sao chép mã từ PDF.')
h('8 1 Các bước thực hiện', 2)
for text in ['1. Giải nén gói bài làm và mở cửa sổ dòng lệnh tại thư mục chứa bayes_server.py.',
             '2. Chạy lệnh python bayes_server.py. Mã tạo dữ liệu Train/Test, đọc lại dữ liệu, học tham số từ Train và đánh giá trên Test.',
             '3. Kiểm tra thông báo “Tat ca kiem tra deu dat” và các tệp trong thư mục ket_qua.',
             '4. Để vẽ lại ba biểu đồ, cài Pillow bằng python -m pip install Pillow, sau đó chạy python draw_charts.py. Phần tính toán Bayes không cần Pillow.']:
    p(text)
h('8 2 Đầu vào và đầu ra', 2)
p('TRAIN và TEST ở đầu mã là các tần số trong Bảng 1. Mỗi dòng train.csv hoặc test.csv gồm id, x_bin và label. id chỉ là số thứ tự trong từng tập, không phải mã máy; label = 1 là Failure. Các mẫu lặp cùng bin và nhãn là cách biểu diễn tần số, không phải các hồ sơ telemetry đã quan sát.')
cap('Bảng 8. Tệp đầu ra và vai trò kiểm chứng')
table(['Tệp trong ket_qua','Nội dung'],[['train.csv và test.csv','4.800 và 1.200 mẫu mô phỏng'],['results.json','Tham số, posterior, chỉ số và khảo sát'],['posterior.csv','Posterior của từng bin'],['cost_sweep.csv','Các kịch bản tỷ lệ chi phí'],['prior_sweep.csv','Các kịch bản prior tại x = 2'],['likelihood.png, cost.png, prior.png','Ba biểu đồ do draw_charts.py tạo']], [7,9.7])
h('8 3 Kết quả kiểm tra', 2)
p('Chương trình kiểm tra tổng số mẫu, số mẫu lỗi, tổng likelihood bằng 1, posterior thuộc [0,1], hai ma trận nhầm lẫn, tổng chi phí 506 và 266, cùng điểm posterior bằng 0,5 tại prior đổi quyết định. Các kiểm tra đã đạt trong lần chạy dùng để tạo báo cáo.')
p('Các posterior chưa làm tròn được sử dụng khi ra quyết định; chỉ làm tròn khi trình bày. Khi sửa dữ liệu hoặc giả định để khảo sát thêm, cần cập nhật các kiểm tra gắn với kịch bản gốc. Mã vẽ đọc results.json thay vì nhập lại số liệu thủ công.')

lines = (BASE/'bayes_server.py').read_text(encoding='utf-8').splitlines()
# Chia tại ranh giới hàm hoặc khối mã để người đọc theo dõi được.
cut1 = lines.index('def fit(samples, alpha=1.0):')
cut2 = lines.index('def main():')
chunks = [lines[:cut1], lines[cut1:cut2], lines[cut2:]]
for idx, lineset in enumerate(chunks):
    page('Phụ lục A Mã Python đầy đủ' if idx == 0 else f'Phụ lục A Mã Python phần {idx+1}')
    if idx == 0:
        p('Tệp bayes_server.py. Ghép liên tiếp ba phần dưới đây để có toàn bộ chương trình; mã không có dấu ba chấm hoặc phần tính toán bị lược bỏ.')
    for line in lineset:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.line_spacing = 1.05
        run = para.add_run(line if line else ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(9)

for root in [doc.styles.element, doc.element]:
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(BASE/'B26CHKH068_Luong_Duc_Manh_BT1_HoanChinh.docx')
print('DOCX saved', len(chunks), 'code pages', len(lines), 'code lines')
